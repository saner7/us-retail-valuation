"""Merger model on small, hand-checkable numbers.

Buyer: price 100, 10 shares, EPS 5 (net income 50), cash 150 (minimum 10% of revenue 1,000 = 100,
so 50 is spare). Target: price 50, 4 shares, EPS 4 (net income 16), book equity 100.
Offer at a 25% premium: 62.5 per share, 250 in total. Tax 25%, new debt 8%, cash earns 4%.
"""
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

import export
import merger
from fake_data import make_data

ACQ = {"ticker": "BUY", "price": 100.0, "shares": 10.0, "eps": 5.0, "growth_y2": 0.0, "revenue": 1000.0,
       "ebitda": 120.0, "debt": 200.0, "cash": 150.0, "net_debt": 50.0, "book_equity": 300.0, "interest": 10.0,
       "tax_rate": 0.25, "fiscal_year_end": "31 Jan"}
TGT = {"ticker": "TGT", "price": 50.0, "shares": 4.0, "eps": 4.0, "growth_y2": 0.0, "revenue": 400.0,
       "ebitda": 40.0, "debt": 50.0, "cash": 0.0, "net_debt": 50.0, "book_equity": 100.0, "interest": 4.0,
       "tax_rate": 0.25, "fiscal_year_end": "31 Jan"}


def inputs(**kw):
    base = {"premium": 0.25, "offer_price": None, "stock_pct": 0.5, "debt_rate": 0.08, "cash_rate": 0.04,
            "refinance_target_debt": False, "cost_synergies": 0.0, "revenue_synergies": 0.0,
            "phase_in": (0.5, 1.0), "integration_cost_x": 1.0, "fee_pct": 0.0, "intangibles_pct": 0.30,
            "amortisation_years": 10, "min_cash_pct": 0.10, "tax_rate": 0.25}
    base.update(kw)
    return base


def test_all_stock():
    r = merger.run(ACQ, TGT, inputs(stock_pct=1.0))
    y = r["years"][0]
    assert r["offer_price"] == 62.5 and r["new_shares"] == pytest.approx(2.5)    # 250 / 100
    assert r["ownership"]["buyer"] == pytest.approx(0.8) and r["ownership"]["target"] == pytest.approx(0.2)
    # PPA: 250 - book 100 = 150 excess; 30% = 45 intangibles over 10 years = 4.5 a year
    assert r["ppa"]["intangibles"] == pytest.approx(45) and y["amortisation"] == pytest.approx(4.5)
    assert y["eps_cash"] == pytest.approx(66 / 12.5)                                # (50 + 16) / 12.5 = 5.28
    assert y["eps_gaap"] == pytest.approx((66 - 4.5 * 0.75) / 12.5)                 # 5.01
    assert y["accretion_gaap"] == pytest.approx(0.01)
    qc = merger.quick_check(ACQ, TGT, inputs(stock_pct=1.0), r)
    assert qc["buyer_pe"] == 20 and qc["offer_pe"] == pytest.approx(15.625) and qc["rule_accretive"]
    assert qc["consistent"]


def test_all_cash_uses_spare_cash_first_then_debt():
    r = merger.run(ACQ, TGT, inputs(stock_pct=0.0))
    su, y = r["sources_uses"]["sources"], r["years"][0]
    assert su["Buyer's excess cash"] == pytest.approx(50) and su["New debt"] == pytest.approx(200)
    # interest 8% x 200 = 16, lost 4% x 50 = 2 -> after tax 13.5; 66 - 13.5 = 52.5 over 10 shares
    assert y["eps_cash"] == pytest.approx(5.25) and y["eps_gaap"] == pytest.approx((52.5 - 3.375) / 10)
    assert r["credit"]["debt"] == pytest.approx(450) and r["credit"]["debt_to_ebitda"] == pytest.approx(450 / 160)
    # target earnings yield 4 / 62.5 = 6.4% vs after-tax cost (50 x 4% + 200 x 8%) / 250 x 0.75 = 5.4%
    qc = merger.quick_check(ACQ, TGT, inputs(stock_pct=0.0), r)
    assert qc["target_earnings_yield"] == pytest.approx(0.064)
    assert qc["cash_cost_after_tax"] == pytest.approx(0.054) and qc["rule_accretive"] and qc["consistent"]


def test_mixed_payment_and_breakevens():
    r = merger.run(ACQ, TGT, inputs())
    y = r["years"][0]
    # 125 in shares (1.25 new), 125 in cash: 50 spare cash + 75 debt; 6 + 2 = 8 x 0.75 = 6
    assert r["pro_forma_shares"] == pytest.approx(11.25) and y["eps_cash"] == pytest.approx(60 / 11.25)
    cash_deal = inputs(stock_pct=0.0)
    y1 = merger.run(ACQ, TGT, cash_deal)["years"][0]
    # GAAP 49.125 vs 50 needed: (50 - 49.125) / 0.75 = 1.1667 realised in Year 1, 2.333 run rate (50% phase-in)
    assert y1["breakeven_synergies"] == pytest.approx(0.875 / 0.75)
    with_syn = merger.run(ACQ, TGT, {**cash_deal, "cost_synergies": y1["breakeven_run_rate"]})
    assert with_syn["years"][0]["accretion_gaap"] == pytest.approx(0, abs=1e-12)
    p = merger.breakeven_premium(ACQ, TGT, inputs())
    at_p = merger.run(ACQ, TGT, inputs(premium=p))
    assert at_p["years"][0]["accretion_gaap"] == pytest.approx(0, abs=1e-9)


def test_refinance_synergy_phase_in_and_one_offs():
    r = merger.run(ACQ, TGT, inputs(stock_pct=0.0, refinance_target_debt=True, cost_synergies=8.0, fee_pct=0.02))
    su = r["sources_uses"]
    # fees 2% x (250 + 50) = 6; cash needed 250 + 6 + 50 = 306: 50 spare cash, 256 debt
    assert su["uses"]["Transaction fees"] == pytest.approx(6) and su["sources"]["New debt"] == pytest.approx(256)
    y1, y2 = r["years"]
    assert y1["target_interest_saved"] == 4 and y1["synergies"] == 4 and y2["synergies"] == 8
    assert r["credit"]["debt"] == pytest.approx(200 + 256)                  # target debt repaid, not added
    assert r["one_off"] == {"transaction_fees": pytest.approx(6), "integration_costs": pytest.approx(8)}


def test_warnings():
    heavy = {**ACQ, "debt": 600.0}
    w = " ".join(merger.run(heavy, TGT, inputs(stock_pct=0.0))["warnings"])       # 850 / 160 = 5.3x
    assert "above 4x" in w and "borrowing spread would likely be higher" in w
    w = " ".join(merger.run(ACQ, {**TGT, "eps": -1.0, "fiscal_year_end": "31 Dec"}, inputs())["warnings"])
    assert "negative earnings" in w and "not calendarised" in w
    assert merger.run(ACQ, TGT, inputs())["warnings"] == []


def test_both_companies_use_the_same_earnings_basis():
    data = make_data(["WMT", "COST"])
    acq, tgt = merger.pair_inputs(data, "WMT", "COST")
    assert acq["eps_source"].startswith("TTM") and tgt["eps_source"].startswith("TTM")
    data["companies"]["WMT"]["meta"]["forward_eps"] = 3.0          # only one has a forecast: still TTM
    assert merger.pair_inputs(data, "WMT", "COST")[1]["eps_source"].startswith("TTM")
    data["companies"]["COST"]["meta"]["forward_eps"] = 4.0
    acq, tgt = merger.pair_inputs(data, "WMT", "COST")
    assert acq["eps"] == 3.0 and tgt["eps"] == 4.0 and "analyst consensus" in tgt["eps_source"]


def test_bridge_adds_up_to_pro_forma_eps_for_both_definitions():
    inp = inputs(refinance_target_debt=True, cost_synergies=8.0)
    r = merger.run(ACQ, TGT, inp)
    for year in (1, 2):
        for basis, key in (("gaap", "eps_gaap"), ("cash", "eps_cash")):
            b = merger.bridge(r, ACQ, year, basis)
            changes = b[b["kind"] == "change"]["value"].sum()
            assert b.iloc[0]["value"] + changes == pytest.approx(b.iloc[-1]["value"])
            assert b.iloc[-1]["value"] == pytest.approx(r["years"][year - 1][key])
    steps = list(merger.bridge(r, ACQ, 1, "gaap")["step"])
    assert "Amortisation of intangibles" in steps and "Target interest saved" in steps
    assert "Amortisation of intangibles" not in list(merger.bridge(r, ACQ, 1, "cash")["step"])


def test_sensitivity_grids_contain_the_users_case():
    inp = inputs(premium=0.25, stock_pct=0.5, cost_synergies=4.0)          # 1% of the target's sales (400)
    t1 = merger.sensitivity_premium_stock(ACQ, TGT, inp)
    assert 0.25 in t1.index and 0.5 in t1.columns
    assert t1.loc[0.25, 0.5] == pytest.approx(merger.run(ACQ, TGT, inp)["years"][0]["accretion_gaap_pct"])
    assert (t1[0.5].diff().dropna() < 0).all()                             # a higher premium never helps
    t2 = merger.sensitivity_synergy_premium(ACQ, TGT, inp)
    assert 0.01 in t2.index and 0.25 in t2.columns
    assert t2.loc[0.01, 0.25] == pytest.approx(t1.loc[0.25, 0.5])
    assert (t2[0.25].diff().dropna() > 0).all()                            # more synergies always help
    cash = merger.sensitivity_premium_stock(ACQ, TGT, inp, "cash")
    assert (cash.values >= t1.values - 1e-12).all()                        # cash EPS leaves out the amortisation


def test_missing_inputs_give_na_instead_of_a_guess():
    bad = {**TGT, "net_debt": float("nan")}                                # e.g. no lease data for the target
    assert merger.missing_inputs(ACQ, bad, inputs()) == ["target net debt"]
    assert merger.missing_inputs(ACQ, TGT, inputs(cash_rate=float("nan"))) == ["cash interest rate"]
    assert merger.missing_inputs(ACQ, TGT, inputs()) == []
    import math
    assert math.isnan(merger.breakeven_premium(ACQ, bad, inputs()))


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
@pytest.mark.parametrize("changes", [{}, {"refinance_target_debt": True, "stock_pct": 0.2, "premium": 0.4}])
def test_excel_merger_sheet_matches_python(tmp_path, changes):
    acq, tgt = {**ACQ, "name": "Buyer", "eps_source": "test"}, {**TGT, "name": "Target", "eps_source": "test"}
    inp = inputs(cost_synergies=12.0, revenue_synergies=3.0, integration_cost_x=1.5, fee_pct=0.02, **changes)
    r = merger.run(acq, tgt, inp)
    src = tmp_path / "ma.xlsx"
    src.write_bytes(export.build_merger_workbook(acq, tgt, inp, r, merger.breakeven_premium(acq, tgt, inp)))
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "out"), str(src)],
                   check=True, capture_output=True, timeout=180)
    ws = load_workbook(tmp_path / "out" / "ma.xlsx", data_only=True)["M&A"]
    y1, y2 = r["years"]
    for cell, expected in [("B86", y1["accretion_gaap_pct"]), ("C86", y2["accretion_gaap_pct"]),
                           ("B88", y1["accretion_cash_pct"]), ("C88", y2["accretion_cash_pct"]),
                           ("B83", y1["eps_gaap"]), ("C84", y2["eps_cash"]), ("B62", r["ownership"]["buyer"]),
                           ("B63", r["ownership"]["target"]), ("B96", r["credit"]["debt_to_ebitda"]),
                           ("B97", r["credit"]["ebitda_to_interest"]), ("B60", r["offer_pe"]),
                           # money cells are in millions in the sheet; the test inputs are tiny, so scale by 1e-6
                           ("B89", y1["breakeven_synergies"] / 1e6), ("C90", y2["breakeven_run_rate"] / 1e6),
                           ("B52", r["sources_uses"]["sources"]["New debt"] / 1e6),
                           ("B53", r["new_shares"] / 1e6), ("B58", r["ppa"]["annual_amortisation"] / 1e6)]:
        assert ws[cell].value == pytest.approx(expected, rel=1e-9), cell


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_excel_matches_python_with_direct_inputs(tmp_path):
    acq, tgt = {**ACQ, "name": "Buyer", "eps_source": "test"}, {**TGT, "name": "Target", "eps_source": "test"}
    inp = inputs(stock_pct=0.0, cost_synergies=6.0, new_debt=120.0, fees_amount=2.5, amortisation_per_year=0.4)
    r = merger.run(acq, tgt, inp)
    src = tmp_path / "ma.xlsx"
    src.write_bytes(export.build_merger_workbook(acq, tgt, inp, r, merger.breakeven_premium(acq, tgt, inp)))
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "out"), str(src)],
                   check=True, capture_output=True, timeout=180)
    ws = load_workbook(tmp_path / "out" / "ma.xlsx", data_only=True)["M&A"]
    y1 = r["years"][0]
    for cell, expected in [("B52", 120.0 / 1e6), ("B46", 2.5 / 1e6), ("B58", 0.4 / 1e6),
                           ("B51", r["sources_uses"]["sources"]["Buyer's excess cash"] / 1e6),
                           ("B86", y1["accretion_gaap_pct"]), ("B88", y1["accretion_cash_pct"])]:
        assert ws[cell].value == pytest.approx(expected, rel=1e-9), cell


def test_direct_inputs_override_the_model_s_own_working():
    r = merger.run(ACQ, TGT, inputs(stock_pct=0.0, new_debt=1000.0, fees_amount=3.0, amortisation_per_year=0.5))
    su = r["sources_uses"]
    # cash needed = 250 + 3; new debt cannot exceed it, and the buyer's cash covers the rest (here: nothing)
    assert su["sources"]["New debt"] == pytest.approx(253) and su["sources"]["Buyer's excess cash"] == pytest.approx(0)
    assert su["uses"]["Transaction fees"] == 3.0 and r["ppa"]["annual_amortisation"] == 0.5
    assert r["ppa"]["intangibles"] != r["ppa"]["intangibles"]            # NaN: amortisation was given directly
    part = merger.run(ACQ, TGT, inputs(stock_pct=0.0, new_debt=200.0))["sources_uses"]["sources"]
    assert part["New debt"] == 200 and part["Buyer's excess cash"] == pytest.approx(50)   # 250 + fees 0 - 200


DKS_FL = "DKS-FL-2025"


def test_example_file_has_a_source_for_every_number():
    df = merger.load_examples()
    assert merger.example_names(df) == [DKS_FL] and len(df) > 50
    assert (df["source"].str.len() > 20).all() and (df["unit"] != "").all()
    numeric = df[df["unit"] != "text"]
    assert numeric["source"].str.contains("sec.gov").all()                 # every figure points to a filing
    assert df[df["note"].str.startswith("DERIVED")]["field"].tolist() == ["t_ebitda_adjusted", "t_impairment_q"]


def test_example_trailing_twelve_months_and_inputs():
    df = merger.load_examples()
    acq, tgt, inp, f = merger.example_pair(df, DKS_FL, "ttm", short_rate=0.04)
    assert acq["eps"] * acq["shares"] == pytest.approx(1165.308e6 - 275.295e6 + 264.288e6)    # 10-K + Q1 - Q1 a year ago
    assert tgt["eps"] * tgt["shares"] == pytest.approx(12e6 - 8e6 - 363e6) and tgt["eps"] < 0
    assert tgt["revenue"] == pytest.approx(7988e6 - 1879e6 + 1794e6)
    assert acq["ebitda"] == pytest.approx(1473.932e6 - 330.799e6 + 366.116e6 + 400.409e6 - 91.477e6 + 97.860e6)
    assert inp["offer_price"] == 24.0 and inp["stock_pct"] == 0 and inp["new_debt"] == 1732e6
    assert inp["debt_rate"] == pytest.approx(28.8e6 / 0.26 / 1732e6) and inp["cost_synergies"] == pytest.approx(112.5e6)
    r = merger.run(acq, tgt, inp)
    assert r["premium"] == pytest.approx(24 / 13.31 - 1)                    # the proxy's 80.3% premium
    assert r["sources_uses"]["sources"]["Buyer's excess cash"] == pytest.approx(95.278e6 * 24 + 56.4e6 - 1732e6)
    assert r["ppa"]["annual_amortisation"] == 885e3


def test_example_results_by_earnings_basis():
    df = merger.load_examples()
    results = {}
    for basis in ("ttm", "mgmt", "mgmt_downside"):
        acq, tgt, inp, f = merger.example_pair(df, DKS_FL, basis, short_rate=0.04)
        results[basis] = merger.run(acq, tgt, inp)["years"]
    assert results["ttm"][0]["accretion_gaap_pct"] < -0.3                      # Foot Locker's TTM is a loss
    assert 0.15 < results["mgmt_downside"][0]["accretion_gaap_pct"] < results["mgmt"][0]["accretion_gaap_pct"]
    acq, tgt, inp, f = merger.example_pair(df, DKS_FL, "mgmt", 0.04)
    assert tgt["eps"] == pytest.approx((591e6 - 202e6 - 17e6) * 0.74 / 95.278e6)   # (EBITDA - D&A - interest)(1 - tax)
    assert tgt["growth_y2"] == pytest.approx((702 - 202 - 17) / (591 - 202 - 17) - 1)
    # mechanics check against Dick's own pro forma on fiscal 2024 actuals (no synergies)
    acq, tgt, inp, f = merger.example_pair(df, DKS_FL, "fy2024", 0.04)
    assert acq["eps"] == pytest.approx(14.05, abs=0.005)                       # matches the 10-K
    check = merger.run(acq, tgt, inp, synergies_on=False)["years"][0]["accretion_gaap_pct"]
    assert -0.09 < check < -0.06 and f["ref_pf_eps_fy2024_cash"] / f["ref_eps_fy2024_standalone"] - 1 < check
