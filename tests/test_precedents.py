"""Stage C: precedent transactions. Small hand-checkable numbers, offline."""

import io
import math

import pandas as pd
import pytest
from openpyxl import load_workbook

import config
import export
import precedents
import summary
from fake_data import make_data


def deals_csv(rows):
    head = ",".join(precedents.COLUMNS)
    lines = [head] + [",".join(str(r.get(c, "")) for c in precedents.COLUMNS) for r in rows]
    return ("\n".join(lines) + "\n").encode()


def deal(target, ev_ebitda, status="completed", tier="core", sector="us_retail_grocery"):
    return {"announcement_date": "2020-01-01", "target": target, "acquirer": "Buyer", "status": status,
            "deal_type": "strategic", "enterprise_value": 1000, "currency": "USD", "ev_ebitda": ev_ebitda,
            "premium_pct": 20, "source_url": "https://example.com", "notes": "test", "tier": tier, "sector": sector}


SAMPLE = deals_csv([deal("A", 6.0), deal("B", 8.0), deal("C", 10.0), deal("D", 12.0),
                    deal("E", 5.0, status="terminated"), deal("F", 20.0, tier="secondary"),
                    deal("G", "", sector="us_retail_department")])


def test_built_in_file_loads_with_sources():
    df = precedents.load_precedents()
    assert len(df) == 7 and df["source_url"].str.startswith("https://www.sec.gov/").all()
    default = df[precedents.included_mask(df)]
    assert sorted(default["target"]) == ["Nordstrom", "Smart & Final Stores", "The Fresh Market", "Whole Foods Market"]
    assert precedents.multiple_stats(default)["median"] == pytest.approx(6.8)    # (6.5 + 7.1) / 2
    assert df["ev_revenue"].isna().all()                                         # EV/Sales is not used


def test_filters_default_and_optional():
    df = precedents.load_precedents(SAMPLE)
    assert list(df[precedents.included_mask(df)]["target"]) == ["A", "B", "C", "D"]   # G has no multiple
    assert "E" in set(df[precedents.included_mask(df, include_terminated=True)]["target"])
    assert "F" in set(df[precedents.included_mask(df, include_secondary=True)]["target"])


def test_stats_range_and_bridge():
    df = precedents.load_precedents(SAMPLE)
    deals = df[precedents.included_mask(df)]
    target = {"ebitda_ttm": 100.0, "net_debt": 200.0, "shares_outstanding": 10.0}
    v = precedents.valuation(target, deals)
    # 6, 8, 10, 12: p25 7.5, median 9, p75 10.5 (linear interpolation, like Excel PERCENTILE.INC)
    assert v["multiples"] == pytest.approx((7.5, 9.0, 10.5))
    # (100 x 9 - 200) / 10 = 70 per share
    assert v["per_share"][1] == pytest.approx(70.0)
    assert v["small_sample"]                                     # 4 deals < 5
    three = precedents.valuation(target, deals.iloc[:3])
    assert three["multiples"] == pytest.approx((6.0, 8.0, 10.0)) and "min-max" in three["method"]
    # the range description follows the rule actually used (quartiles from 4 deals)
    assert v["method"].startswith("25th-75th percentile (4")
    assert precedents.football_field_row(target, deals)["basis"].startswith("25th-75th percentile (4 deals)")
    assert precedents.football_field_row(target, deals.iloc[:3])["basis"].startswith("min-max (only 3 deals)")


def test_file_errors_are_explained():
    with pytest.raises(precedents.PrecedentFileError, match="Missing columns"):
        precedents.load_precedents(b"target,ev_ebitda\nA,5\n")
    bad = deals_csv([deal("A", "lots")])
    with pytest.raises(precedents.PrecedentFileError, match="ev_ebitda must hold numbers"):
        precedents.load_precedents(bad)
    with pytest.raises(precedents.PrecedentFileError, match="status"):
        precedents.load_precedents(deals_csv([deal("A", 5, status="maybe")]))
    assert precedents.template_csv().decode().startswith("announcement_date,target")


def test_row_only_for_matching_sub_sector():
    df = precedents.load_precedents()
    deals = df[precedents.included_mask(df)]                    # grocery + department store deals
    assert precedents.applicable({}, "TGT", deals)              # the list was selected for Target
    kroger = {"country": "United States", "industry": "Grocery Stores"}
    home_depot = {"country": "United States", "industry": "Home Improvement Retail"}
    ferrari = {"country": "Italy", "industry": "Auto Manufacturers"}
    assert precedents.applicable(kroger, "KR", deals)
    assert not precedents.applicable(home_depot, "HD", deals)
    assert not precedents.applicable(ferrari, "RACE.MI", deals)
    # With the second-tier specialty deals switched on, Home Depot's sub-sector is covered
    assert precedents.applicable(home_depot, "HD", df[precedents.included_mask(df, include_secondary=True)])


def test_football_field_and_excel_include_the_row():
    data = make_data()
    choices = summary.default_choices(data)
    kw = summary.build_model(data, choices)["run_kwargs"]
    df = precedents.load_precedents()
    row = precedents.football_field_row(data["companies"][config.TARGET]["values"], df[precedents.included_mask(df)])
    ff = summary.football_field_data(data["companies"], kw, precedent_row=row)
    keys = list(ff["key"])
    assert keys.index("prec_ev_ebitda") == keys.index("ev_sales") + 1      # below the comps rows
    assert "prec_ev_ebitda" not in set(summary.football_field_data(data["companies"], kw)["key"])
    wb = load_workbook(io.BytesIO(export.build_workbook(data, choices)))
    texts = [c.value for r in wb["Precedents"].iter_rows() for c in r if isinstance(c.value, str)]
    assert "Whole Foods Market" in texts and any("Small sample" in t for t in texts)
    assert "Precedents: EV / EBITDA" in [c.value for c in wb["Summary"]["A"]]
