"""Stage B: scenario analysis. Small hand-checkable numbers, offline."""

import io
import math
import shutil
import subprocess

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

import config
import dcf
import export
import scenarios
import summary
from fake_data import make_data


def test_halfway_rule_is_symmetric_around_base():
    hist = pd.Series([0.06, 0.08, 0.10])
    # Base 9%: furthest year is 6% (3 points away) -> halfway = 1.5 points each side
    assert scenarios._toward(0.09, hist, True, 0.5) == pytest.approx((0.075, 0.105))
    # Capex: higher is worse, so Bear gets the higher value
    assert scenarios._toward(0.09, hist, False, 0.5) == pytest.approx((0.105, 0.075))
    # No history: Bear = Bull = Base
    assert scenarios._toward(0.09, pd.Series(dtype=float), True, 0.5) == (0.09, 0.09)


@pytest.fixture
def data():
    d = make_data()
    # Give the target a history with spread: margins 6%, 8%, 10%, 8%
    h = d["companies"][config.TARGET]["history"]
    h["ebitda"] = h["revenue"] * np.array([0.06, 0.08, 0.10, 0.08])
    return d


def test_default_scenarios_follow_the_rules(data):
    base = summary.default_choices(data)
    sc = scenarios.default_scenarios(data, config.TARGET, config.PEERS, base)
    assert sc["base"]["margin_end"] == pytest.approx(base["margin_end"])
    # Margin end: halfway to the furthest historical margin, both directions
    h = dcf.historical_ratios(data["companies"][config.TARGET]["history"])["ebitda_margin"]
    step = 0.5 * max(base["margin_end"] - h.min(), h.max() - base["margin_end"])
    assert sc["bear"]["margin_end"] == pytest.approx(base["margin_end"] - step)
    assert sc["bull"]["margin_end"] == pytest.approx(base["margin_end"] + step)
    # Same starting margin in every scenario
    assert sc["bear"]["margin_start"] == sc["bull"]["margin_start"] == base["margin_start"]
    # Terminal growth -/+ 0.5 points, Year 5 growth follows it
    assert sc["bear"]["terminal_growth"] == pytest.approx(base["terminal_growth"] - 0.005)
    assert sc["bull"]["growth_end"] == pytest.approx(base["terminal_growth"] + 0.005)
    assert all(f in sc["rationale"] for f in scenarios.FIELDS)


def test_bull_exit_multiple_cap(data):
    base = summary.default_choices(data)
    sc = scenarios.default_scenarios(data, config.TARGET, config.PEERS, base)
    # fake peers' EV/EBITDA: 12, 8, 14, 6, 9 -> p25 8, median 9, p75 12; target own ~10
    own = data["companies"][config.TARGET]["values"]["enterprise_value"] / data["companies"][config.TARGET]["values"]["ebitda_ttm"]
    assert sc["bear"]["exit_multiple"] == pytest.approx((8 + 9) / 2)
    assert sc["bull_exit_uncapped"] == pytest.approx((9 + 12) / 2)
    assert sc["bull_exit_cap"] == pytest.approx(1.35 * own)
    assert sc["bull"]["exit_multiple"] == pytest.approx(min(10.5, 1.35 * own))
    # A cap below Base never pushes Bull under Base
    cheap = make_data()
    cheap["companies"][config.TARGET]["values"]["enterprise_value"] *= 0.3     # target at ~3x EBITDA
    b = summary.default_choices(cheap)
    s2 = scenarios.default_scenarios(cheap, config.TARGET, config.PEERS, b)
    assert s2["bull"]["exit_multiple"] == pytest.approx(b["exit_multiple"])


def test_weighted_value_needs_weights_of_100_percent():
    values = {"bear": 50.0, "base": 100.0, "bull": 200.0}
    assert scenarios.weighted_value(values, {"bear": 0.25, "base": 0.5, "bull": 0.25}) == pytest.approx(112.5)
    assert math.isnan(scenarios.weighted_value(values, {"bear": 0.3, "base": 0.5, "bull": 0.3}))   # 110%
    assert math.isnan(scenarios.weighted_value({**values, "bear": np.nan}, config.SCENARIO_WEIGHTS))


def test_year_by_year_edits_for_base_only_touch_shared_rows_elsewhere():
    base = {"per_year": {"capex_pct": [0.1] * 5, "tax_rate": [0.3] * 5}, "capex_pct": 0.03}
    bear = scenarios.scenario_choices(base, {"capex_pct": 0.05}, is_base=False)
    assert bear["per_year"] == {"tax_rate": [0.3] * 5}          # scenario capex wins, shared tax stays
    assert scenarios.scenario_choices(base, {}, is_base=True)["per_year"] == base["per_year"]


def test_scenarios_rank_and_feed_the_football_field(data):
    base = summary.default_choices(data)
    sc = scenarios.default_scenarios(data, config.TARGET, config.PEERS, base)
    runs = scenarios.run_scenarios(data, base, sc, config.TARGET)
    table = scenarios.results_table(runs, config.SCENARIO_WEIGHTS)
    assert table.loc["gordon_value", "bear"] < table.loc["gordon_value", "base"] < table.loc["gordon_value", "bull"]
    assert table.loc["gordon_value", "base"] == pytest.approx(dcf.run_dcf(**summary.build_model(data, base)["run_kwargs"])["gordon"]["value_per_share"])
    ff = summary.football_field_data(data["companies"], runs["base"]["run_kwargs"], scenario_runs=runs).set_index("key")
    assert ff.loc["dcf_gordon", "bear"] == pytest.approx(table.loc["gordon_value", "bear"])
    assert ff.loc["dcf_exit", "bull"] == pytest.approx(table.loc["exit_value", "bull"])
    assert math.isnan(ff.loc["ev_ebitda", "bear"])                 # markers only on the DCF rows


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_excel_switch_and_side_by_side_results_match_python(data, tmp_path):
    base = summary.default_choices(data)
    sc = scenarios.default_scenarios(data, config.TARGET, config.PEERS, base)
    runs = scenarios.run_scenarios(data, base, sc, config.TARGET)
    table = scenarios.results_table(runs, config.SCENARIO_WEIGHTS)
    for switch, name in [(1, "bear"), (2, "base"), (3, "bull")]:
        raw = export.build_workbook(data, base, scenario_values={s: sc[s] for s in config.SCENARIOS},
                                    scenario_switch=switch)
        src = tmp_path / f"s{switch}.xlsx"
        src.write_bytes(raw)
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "out"), str(src)],
                       check=True, capture_output=True, timeout=180)
        wb = load_workbook(tmp_path / "out" / f"s{switch}.xlsx", data_only=True)
        # The DCF sheet follows the switch...
        assert wb["DCF"]["B46"].value == pytest.approx(table.loc["gordon_value", name], rel=1e-9)
        assert wb["DCF"]["C46"].value == pytest.approx(table.loc["exit_value", name], rel=1e-9)
        # ...and the Scenarios sheet always values all three, plus the weighted value
        res = wb["Scenarios"]
        for col, s in zip("BCD", config.SCENARIOS):
            assert res[f"{col}27"].value == pytest.approx(table.loc["gordon_value", s], rel=1e-9)
            assert res[f"{col}29"].value == pytest.approx(table.loc["exit_value", s], rel=1e-9)
        assert res["E27"].value == pytest.approx(table.loc["gordon_value", "weighted"], rel=1e-9)
        assert res["E29"].value == pytest.approx(table.loc["exit_value", "weighted"], rel=1e-9)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_excel_weighted_value_is_na_when_weights_do_not_add_up(data, tmp_path):
    base = summary.default_choices(data)
    raw = export.build_workbook(data, base, weights={"bear": 0.3, "base": 0.5, "bull": 0.3})
    (tmp_path / "w.xlsx").write_bytes(raw)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "out"),
                    str(tmp_path / "w.xlsx")], check=True, capture_output=True, timeout=180)
    wb = load_workbook(tmp_path / "out" / "w.xlsx", data_only=True)
    assert wb["Scenarios"]["E27"].value == "#N/A"
    assert wb["Scenarios"]["B23"].value == pytest.approx(1.1)
