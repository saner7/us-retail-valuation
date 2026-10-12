"""Tests for export.py. The structure tests run everywhere; the formula test
needs LibreOffice (to calculate the formulas) and is skipped without it."""

import io
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

import dcf
import export
import summary
from fake_data import make_data

SHEETS = ["Summary", "DCF", "Scenarios", "WACC", "Sensitivity", "Comps", "Bridge", "Precedents", "Market data",
          "Data quality", "History"]


@pytest.fixture(scope="module")
def data():
    return make_data()


def workbook(data, **changes):
    choices = {**summary.default_choices(data), **changes}
    return export.build_workbook(data, choices), choices


def test_sheets_and_which_ones_have_formulas(data):
    raw, _ = workbook(data)
    wb = load_workbook(io.BytesIO(raw))
    assert wb.sheetnames == SHEETS
    for ws in wb.worksheets:
        formulas = [c for row in ws.iter_rows() for c in row if c.data_type == "f"]
        if ws.title in ("DCF", "WACC", "Scenarios"):
            assert formulas, f"{ws.title} should use live formulas"
        else:
            assert not formulas, f"{ws.title} should hold values only"
    assert wb["DCF"]["B12"].value == "=WACC!$C$18"            # DCF discounts at the WACC sheet's result
    assert wb["DCF"]["A42"].value == "= Enterprise value"      # a label, not a formula
    assert wb["DCF"]["A42"].data_type == "s"


def test_inputs_are_written_as_numbers_in_millions(data):
    raw, _ = workbook(data)
    wb = load_workbook(io.BytesIO(raw))
    target = data["companies"]["TGT"]["values"]
    assert wb["DCF"]["B5"].value == pytest.approx(target["revenue_ttm"] / 1e6)
    assert wb["WACC"]["C5"].value == pytest.approx(0.04)
    comps_text = [c.value for row in wb["Comps"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any("not meaningful for retailers" in t for t in comps_text)   # P/B note


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
@pytest.mark.parametrize("mid_year", [True, False])
def test_excel_formulas_match_python(data, tmp_path, mid_year):
    raw, choices = workbook(data, mid_year=mid_year,
                            per_year={"capex_pct": [0.05, 0.045, 0.04, 0.035, 0.035]})
    src = tmp_path / "model.xlsx"
    src.write_bytes(raw)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "out"), str(src)],
                   check=True, capture_output=True, timeout=180)
    wb = load_workbook(tmp_path / "out" / "model.xlsx", data_only=True)
    model = summary.build_model(data, choices)
    r = dcf.run_dcf(**model["run_kwargs"])
    assert wb["WACC"]["C18"].value == pytest.approx(model["wacc"]["wacc"], rel=1e-9)
    assert wb["DCF"]["B46"].value == pytest.approx(r["gordon"]["value_per_share"], rel=1e-9)
    assert wb["DCF"]["C46"].value == pytest.approx(r["exit"]["value_per_share"], rel=1e-9)
    assert wb["DCF"]["B50"].value == pytest.approx(r["gordon"]["implied_exit_multiple"], rel=1e-9)
    assert wb["DCF"]["C50"].value == pytest.approx(r["exit"]["implied_growth"], rel=1e-9)


def test_euro_workbook_without_risk_free_rate(tmp_path):
    data = make_data(["RACE.MI", "MC.PA", "RMS.PA"], currencies={t: "EUR" for t in ["RACE.MI", "MC.PA", "RMS.PA"]},
                     risk_free=None)
    raw = export.build_workbook(data, summary.default_choices(data, "RACE.MI"), "RACE.MI", ["MC.PA", "RMS.PA"])
    wb = load_workbook(io.BytesIO(raw))
    assert wb["WACC"]["C5"].value is None                              # left empty, never guessed
    assert "ENTER" in wb["WACC"]["D5"].value and "Bund" in wb["WACC"]["D5"].value
    assert wb["WACC"]["C18"].value.startswith("=IF(ISBLANK(C5),NA()")  # #N/A instead of a 0% rate
    assert "€" in wb["DCF"]["B10"].number_format
    assert "EUR" in wb["DCF"]["A5"].value
    if shutil.which("soffice"):
        src = tmp_path / "eur.xlsx"
        src.write_bytes(raw)
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "out"),
                        str(src)], check=True, capture_output=True, timeout=180)
        calc = load_workbook(tmp_path / "out" / "eur.xlsx", data_only=True)
        assert calc["WACC"]["C18"].value == "#N/A"
        assert calc["DCF"]["B46"].value == "#N/A"                      # value per share: N/A, not a number
